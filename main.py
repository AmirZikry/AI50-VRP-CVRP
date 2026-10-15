import time
import copy
import vrplib
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from Utils.dataLoader import load_instance
from Utils.evaluator import evaluate_solution

# Construction / Metaheuristic Solvers
from Solvers.NearestNeighbor import NearestNeighborSolver
from Solvers.CheapestInsertion import CheapestInsertionSolver
from Solvers.ClarkeWright import ClarkeWrightSolver
from Solvers.GeneticAlgo import GeneticAlgorithmSolver 
from Solvers.AntColony import AntColonySolver

# Post-Optimization Metaheuristics
from Optimizer.SimAnneal import Simulated_Annealing
from Optimizer.TabuSearch import Tabu_Search


def get_solvers(instance):
    """Factory defining all baseline solvers."""
    return {
        "Nearest Neighbor": lambda: NearestNeighborSolver(instance),
        "Cheapest Insertion": lambda: CheapestInsertionSolver(instance),
        "Clarke-Wright": lambda: ClarkeWrightSolver(instance),
        "Genetic Algorithm": lambda: GeneticAlgorithmSolver(instance, pop_size=100, generations=200, seed=42),
        "Ant Colony": lambda: AntColonySolver(instance, num_ants=50, num_iterations=100, alpha=0.95, beta=0.45, rho=0.3, Q=0.5),
    }


def get_optimizers():
    """Factory defining all post-optimizers."""
    return {
        "Simulated Annealing": lambda sol: Simulated_Annealing(sol),
        "Tabu Search": lambda sol: Tabu_Search(sol, tabu_tenure=10, max_iterations=500),
    }


def export_to_excel(results, instance, best_cost, output_file="benchmark_results.xlsx"):
    """Creates a clean, professionally formatted Excel workbook."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Benchmark Results"
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") 
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")   
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10, color="0F172A")
    title_font = Font(name="Segoe UI", size=15, bold=True, color="1E293B")
    subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="64748B")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Title block
    ws.merge_cells("A1:G1")
    ws["A1"] = f"CVRP Benchmark Report: {instance.name}"
    ws["A1"].font = title_font

    ws["A2"] = f"Dimension: {instance.dimension - 1} customers | Vehicle Capacity: {instance.capacity} | Best-Known Cost: {best_cost}"
    ws["A2"].font = subtitle_font

    headers = [
        "Config Type",
        "Base Solver",
        "Optimizer",
        "Total Cost",
        "Optimality Gap",
        "Runtime (s)",
        "Feasible / Valid"
    ]

    # Write headers
    row_idx = 4
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center" if col_idx > 3 else "left", vertical="center")
    ws.row_dimensions[row_idx].height = 26

    # Write data rows
    for item in results:
        row_idx += 1
        row_cells = [
            ws.cell(row=row_idx, column=1, value=item["Config Type"]),
            ws.cell(row=row_idx, column=2, value=item["Base Solver"]),
            ws.cell(row=row_idx, column=3, value=item["Optimizer"]),
            ws.cell(row=row_idx, column=4, value=item["Total Cost"]),
            ws.cell(row=row_idx, column=5, value=item["Optimality Gap"]),
            ws.cell(row=row_idx, column=6, value=item["Runtime (s)"]),
            ws.cell(row=row_idx, column=7, value=item["Feasible / Valid"]),
        ]

        is_even = (row_idx % 2 == 0)
        for c_idx, cell in enumerate(row_cells, 1):
            cell.font = data_font
            cell.border = thin_border
            if is_even:
                cell.fill = zebra_fill

            # Alignments & formats
            if c_idx in (1, 2, 3):
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif c_idx == 4:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"
            elif c_idx == 5:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "+0.00%;-0.00%;0.00%"
            elif c_idx == 6:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.000"
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # Dynamic column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    wb.save(output_file)
    print(f"\n[OK] Results successfully exported to '{output_file}'")

def main():
    instance_path = "./DataCVRP/instances/XML100_1111_01.vrp"
    solution_path = "./DataCVRP/solutions/XML100_1111_01.sol"

    # 1. Load instance and benchmark reference
    instance = load_instance(instance_path, is_local=True)
    best_known = vrplib.read_solution(solution_path)
    best_cost = best_known["cost"]

    print(f"Loaded: {instance.name} | Customers: {instance.dimension - 1} | Capacity: {instance.capacity}")
    print(f"Best Known Cost: {best_cost}\n")

    solver_factories = get_solvers(instance)
    optimizer_factories = get_optimizers()

    results = []
    initial_solutions = {}

    # --- Phase 1: Standalone Solvers ---
    print("Running Standalone Solvers")
    for solver_name, solver_builder in solver_factories.items():
        print(f" Testing {solver_name}...")
        start = time.perf_counter()
        try:
            solver = solver_builder()
            sol = solver.solve()
            runtime = time.perf_counter() - start

            initial_solutions[solver_name] = sol
            eval_res = evaluate_solution(instance, sol, best_known_cost=best_cost)
            cost = getattr(eval_res, "total_cost", getattr(eval_res, "cost", None))
            gap = ((cost - best_cost) / best_cost) if cost else None

            results.append({
                "Config Type": "Standalone",
                "Base Solver": solver_name,
                "Optimizer": "None",
                "Total Cost": cost,
                "Optimality Gap": gap,
                "Runtime (s)": round(runtime, 3),
                "Feasible / Valid": "Yes" if getattr(eval_res, "is_valid", True) else "No",
            })
        except Exception as e:
            results.append({
                "Config Type": "Standalone",
                "Base Solver": solver_name,
                "Optimizer": "None",
                "Total Cost": "Error",
                "Optimality Gap": "N/A",
                "Runtime (s)": round(time.perf_counter() - start, 3),
                "Feasible / Valid": f"Failed: {str(e)[:20]}",
            })

    # --- Phase 2: Hybrid Solutions ---
    print("\nRunning Solver + Optimizer Hybrids")
    for solver_name, init_sol in initial_solutions.items():
        if init_sol is None:
            continue

        for opt_name, opt_builder in optimizer_factories.items():
            print(f"  -> Testing {solver_name} + {opt_name}...")
            start = time.perf_counter()
            try:

                sol_copy = copy.deepcopy(init_sol)
                optimizer = opt_builder(sol_copy)
                opt_sol = optimizer.optimize()
                runtime = time.perf_counter() - start

                eval_res = evaluate_solution(instance, opt_sol, best_known_cost=best_cost)
                cost = getattr(eval_res, "total_cost", getattr(eval_res, "cost", None))
                gap = ((cost - best_cost) / best_cost) if cost else None

                results.append({
                    "Config Type": "Hybrid",
                    "Base Solver": solver_name,
                    "Optimizer": opt_name,
                    "Total Cost": cost,
                    "Optimality Gap": gap,
                    "Runtime (s)": round(runtime, 3),
                    "Feasible / Valid": "Yes" if getattr(eval_res, "is_valid", True) else "No",
                })
            except Exception as e:
                results.append({
                    "Config Type": "Hybrid",
                    "Base Solver": solver_name,
                    "Optimizer": opt_name,
                    "Total Cost": "Error",
                    "Optimality Gap": "N/A",
                    "Runtime (s)": round(time.perf_counter() - start, 3),
                    "Feasible / Valid": f"Failed: {str(e)[:20]}",
                })

    # --- Phase 3: Export to Excel ---
    export_to_excel(results, instance, best_cost, output_file="cvrp_benchmark_results.xlsx")


if __name__ == "__main__":
    main()